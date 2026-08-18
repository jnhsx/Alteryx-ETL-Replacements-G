from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from pathlib import Path
import re

import pandas as pd
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from shared.pipeline import PipelineArtifact, PipelineResult


EXCEL_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HTML_CONTENT_TYPE = "text/html; charset=utf-8"


@dataclass(frozen=True)
class BdxTrackerConfig:
    selected_month: str
    version: str = ""


def run(input_bytes: bytes, filename: str, config: BdxTrackerConfig) -> PipelineResult:
    if not input_bytes:
        raise ValueError("BDX Tracker input blob is empty.")

    selected_month_dt = parse_selected_month(config.selected_month)
    df = pd.read_excel(BytesIO(input_bytes), engine="openpyxl")
    output_df = transform(df, config.selected_month, selected_month_dt)

    source_stem = Path(filename).stem
    safe_month = safe_path_segment(config.selected_month)
    version_label = f" v{safe_path_segment(config.version)}" if config.version.strip() else ""
    excel_filename = f"BDX Management {safe_month}{version_label} ({source_stem}).xlsx"
    html_filename = f"BDX Email {safe_month}{version_label} ({source_stem}).html"

    excel_bytes = build_excel(output_df)
    html_bytes = build_done_html(output_df).encode("utf-8")

    html_artifact = PipelineArtifact(
        output_bytes=html_bytes,
        output_filename=html_filename,
        content_type=HTML_CONTENT_TYPE,
    )

    return PipelineResult(
        output_bytes=excel_bytes,
        output_filename=excel_filename,
        content_type=EXCEL_CONTENT_TYPE,
        row_count=len(output_df),
        artifacts=[html_artifact],
    )


def transform(
    df: pd.DataFrame,
    selected_month: str,
    selected_month_dt: pd.Timestamp,
) -> pd.DataFrame:
    df = df.copy()
    df.columns = df.columns.str.strip()
    validate_required_columns(df)

    month_columns = detect_month_columns(df)
    if selected_month not in month_columns:
        available_months = ", ".join(sorted(month_columns.keys()))
        raise ValueError(
            f"Month '{selected_month}' not found in file. Available months: {available_months}"
        )

    df["Status_raw"] = df[selected_month]
    df["Status_clean"] = (
        df["Status_raw"]
        .astype(str)
        .str.strip()
        .str.lower()
        .str.replace(".", "", regex=False)
    )

    status_mapping = {
        "not received": "Not Received",
        "in progress": "In Progress",
        "partially done": "Partially Done",
        "done": "Done",
        "no longer required": "No Longer Required",
        "on hold": "On Hold",
    }

    df["Status"] = df["Status_clean"].map(status_mapping)
    df["Status"] = df["Status"].fillna(df["Status_raw"])

    due_date_column = next(
        (col for col in df.columns if "due" in col.lower() and "date" in col.lower()),
        None,
    )

    df["Due Date for Submission"] = df.apply(
        lambda row: build_due_date(row, due_date_column, selected_month_dt),
        axis=1,
    )
    df["Submission Status"] = df["Due Date for Submission"].apply(calculate_submission_status)
    df["GWP Clean"] = df["GWP"].apply(clean_currency) if "GWP" in df.columns else None

    output_df = pd.DataFrame(
        {
            "Scheme Name": df["Scheme Name"],
            "Month": selected_month_dt.strftime("%y-%b"),
            "Status": df["Status"],
            "Frequency": df.get("Frequency", ""),
            "Region": df.get("Region", ""),
            "Due Date for Submission": df["Due Date for Submission"],
            "Submission Status": df["Submission Status"],
            "Comment": df.get("Notes", ""),
            "GWP Received on Bdx": df["GWP Clean"],
            "GWP Booked": df["GWP Clean"],
        }
    )

    status_order = [
        "Not Received",
        "In Progress",
        "Partially Done",
        "On Hold",
        "No Longer Required",
        "Done",
    ]
    output_df["Status"] = pd.Categorical(
        output_df["Status"],
        categories=status_order,
        ordered=True,
    )

    return output_df.sort_values(
        by=["Status", "Due Date for Submission"],
        ascending=[True, True],
        na_position="last",
    )


def validate_required_columns(df: pd.DataFrame) -> None:
    required = ["Scheme Name"]
    missing = [col for col in required if col not in df.columns]
    if missing:
        raise ValueError(f"BDX Tracker input is missing required columns: {', '.join(missing)}")


def parse_selected_month(selected_month: str) -> pd.Timestamp:
    try:
        return pd.Timestamp(datetime.strptime(selected_month, "%b-%y"))
    except ValueError as exc:
        raise ValueError(
            f"Invalid selected_month '{selected_month}'. Expected format: Jan-26"
        ) from exc


def detect_month_columns(df: pd.DataFrame) -> dict[str, pd.Timestamp]:
    pattern = re.compile(r"^[A-Za-z]{3}-\d{2}$")
    month_cols = {}
    for col in df.columns:
        col_clean = str(col).strip()
        if pattern.match(col_clean):
            try:
                month_cols[col_clean] = pd.Timestamp(datetime.strptime(col_clean, "%b-%y"))
            except ValueError:
                pass
    return month_cols


def clean_currency(value: object) -> float | None:
    if pd.isna(value) or str(value).strip() == "":
        return None

    cleaned = str(value).replace("£", "").replace("Â£", "").replace(",", "").strip()
    try:
        return float(cleaned)
    except ValueError:
        return None


def build_due_date(
    row: pd.Series,
    due_date_column: str | None,
    selected_month_dt: pd.Timestamp,
) -> pd.Timestamp:
    if not due_date_column:
        return pd.NaT

    due_day = row[due_date_column]
    if pd.isna(due_day):
        return pd.NaT

    try:
        return pd.Timestamp(
            datetime(
                year=selected_month_dt.year,
                month=selected_month_dt.month,
                day=int(due_day),
            )
        )
    except (TypeError, ValueError):
        return pd.NaT


def calculate_submission_status(due_date: pd.Timestamp) -> str:
    if pd.isna(due_date):
        return ""

    today = datetime.today().date()
    diff = (today - due_date.date()).days

    if diff > 0:
        return f"Delayed by {diff} days"
    if diff < 0:
        return f"Due in {abs(diff)} days"
    return "Due Today"


def build_excel(df: pd.DataFrame) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        format_workbook(writer, df)
    return output.getvalue()


def format_workbook(writer: pd.ExcelWriter, df: pd.DataFrame) -> None:
    header_fill = PatternFill(start_color="C0C0C0", end_color="C0C0C0", fill_type="solid")
    header_font = Font(bold=True, color="000000")
    band_fill_1 = PatternFill(start_color="F0F0F0", end_color="F0F0F0", fill_type="solid")
    band_fill_2 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    status_colours = {
        "not received": PatternFill(start_color="FF5555", end_color="FF5555", fill_type="solid"),
        "in progress": PatternFill(start_color="FFC000", end_color="FFC000", fill_type="solid"),
        "partially done": PatternFill(start_color="00B0F0", end_color="00B0F0", fill_type="solid"),
        "done": PatternFill(start_color="00B050", end_color="00B050", fill_type="solid"),
        "no longer required": PatternFill(start_color="000000", end_color="000000", fill_type="solid"),
        "on hold": PatternFill(start_color="7030A0", end_color="7030A0", fill_type="solid"),
    }

    df.to_excel(writer, sheet_name="Overall", index=False)

    used_sheet_names = {"Overall"}
    for region in sorted(df["Region"].dropna().unique()):
        region_df = df[df["Region"] == region]
        sheet_name = unique_sheet_name(str(region), used_sheet_names)
        region_df.to_excel(writer, sheet_name=sheet_name, index=False)

    thin = Side(style="thin")
    cell_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    white_font = Font(color="FFFFFF")

    for sheet_name in writer.book.sheetnames:
        ws = writer.book[sheet_name]
        max_row = ws.max_row
        max_col = ws.max_column
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        headers = [ws.cell(row=1, column=col).value for col in range(1, max_col + 1)]

        for col in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.border = cell_border

        for row in range(2, max_row + 1):
            band_fill = band_fill_1 if row % 2 == 0 else band_fill_2
            for col in range(1, max_col + 1):
                cell = ws.cell(row=row, column=col)
                cell.fill = band_fill
                cell.font = Font(color="000000")
                cell.border = cell_border

                if headers[col - 1] == "Due Date for Submission" and cell.value:
                    cell.number_format = "yyyy-mm-dd"

                if headers[col - 1] in ["GWP Received on Bdx", "GWP Booked"] and cell.value:
                    cell.number_format = "£#,##0.00"

        if "Status" in headers:
            status_col_index = headers.index("Status") + 1
            for row in range(2, max_row + 1):
                cell = ws.cell(row=row, column=status_col_index)
                status_value = str(cell.value).strip().lower()
                if status_value in status_colours:
                    cell.fill = status_colours[status_value]
                    if status_value == "no longer required":
                        cell.font = white_font

        for col in range(1, max_col + 1):
            ws.column_dimensions[get_column_letter(col)].width = 20


def unique_sheet_name(region: str, used_sheet_names: set[str]) -> str:
    invalid_chars = r"[]:*?/\\"
    base = "".join("_" if char in invalid_chars else char for char in region).strip()
    base = base or "Region"
    base = base[:31]

    sheet_name = base
    index = 2
    while sheet_name in used_sheet_names:
        suffix = f"_{index}"
        sheet_name = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1

    used_sheet_names.add(sheet_name)
    return sheet_name


def build_done_html(df: pd.DataFrame) -> str:
    # Calculate GWP from Done status rows for header
    done_df = df[df["Status"] == "Done"].copy()
    total_done_gwp = done_df["GWP Booked"].fillna(0).sum()
    
    # Get rows with Not Received or On Hold status for the table
    table_df = df[df["Status"].isin(["Not Received", "On Hold"])].copy()
    # Sort to show Not Received first, then On Hold
    table_df["Status"] = pd.Categorical(
        table_df["Status"],
        categories=["Not Received", "On Hold"],
        ordered=True
    )
    table_df = table_df.sort_values(by="Status")
    rows_html = ""

    for _, row in table_df.iterrows():
        submission_status = row.get("Submission Status", "")
        colour = "red" if "Delayed" in str(submission_status) else "black"
        status_value = row.get("Status", "")
        status_bg = "#FF5555" if status_value == "Not Received" else "#7030A0"
        rows_html += f"""
        <tr>
            <td align="center">{html_value(row.get("Scheme Name", ""))}</td>
            <td align="center">{html_value(row.get("Month", ""))}</td>
            <td align="center" style="background:{status_bg};">{html_value(status_value)}</td>
            <td align="center">{html_value(row.get("Frequency", ""))}</td>
            <td align="center">{html_value(row.get("Region", ""))}</td>
            <td align="center">{format_date(row.get("Due Date for Submission"))}</td>
            <td align="center" style="color:{colour};">{html_value(submission_status)}</td>
            <td align="center">{html_value(row.get("Comment", ""))}</td>
        </tr>
        """

    return f"""
<html>
<head>
<meta charset="UTF-8">
</head>
<body style="font-family:Arial;">
<p>Find the detailed RM Tracker</p>
<p>Overall DONE GWP booked: {format_currency(total_done_gwp)}</p>
<br>

<table border="1" cellspacing="4" cellpadding="4" width="100%" style="border:solid black 1pt;">
<thead style="background:silver;">
<tr>
<th>Scheme Name</th>
<th>Month</th>
<th>Status</th>
<th>Frequency</th>
<th>Region</th>
<th>Due Date for Submission</th>
<th>Submission Status</th>
<th>Comment</th>
</tr>
</thead>
<tbody>
{rows_html}
</tbody>
</table>
</body>
</html>
"""


def html_value(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value)


def format_currency(value: object) -> str:
    if pd.isna(value):
        return ""
    return f"£{float(value):,.2f}"


def format_date(value: object) -> str:
    if pd.isna(value):
        return ""
    return pd.Timestamp(value).strftime("%Y-%m-%d")


def safe_path_segment(value: str) -> str:
    return value.strip().replace("/", "-").replace("\\", "-")
